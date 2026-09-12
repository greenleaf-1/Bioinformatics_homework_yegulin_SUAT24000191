from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
TEMPLATE = Path("/Users/theaye/.codex/plugins/cache/openai-curated-remote/openai-templates/0.1.1/skills/artifact-template-experiment-analysis/assets/reference.docx")
OUTPUT = ROOT / "Week3_Homework_Ye_Gulin_SUAT24000191.docx"
R_SCRIPT = ROOT / "week3_homework_analysis.R"
GREEN = "137A3D"
LIGHT_GREEN = "EAF4ED"
LIGHT_GRAY = "F1F2F1"


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, value=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def run(paragraph, text, bold=False, italic=False, size=10.5, color=None, font="Georgia"):
    item = paragraph.add_run(text)
    item.bold = bold
    item.italic = italic
    item.font.size = Pt(size)
    item.font.name = font
    item._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    if color:
        item.font.color.rgb = RGBColor.from_string(color)
    return item


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    run(p, text, bold=True, size=22 if level == 1 else 15, color=GREEN)
    return p


def body(doc, text, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.05
    run(p, text, italic=italic)
    return p


def table(doc, headers, rows, widths, font_size=9.1):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    # The reference template does not include Word's built-in Table Grid
    # style, so draw stable borders directly while retaining its typography.
    tbl_borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "A7ACA9")
        tbl_borders.append(border)
    t._tbl.tblPr.append(tbl_borders)
    repeat_header(t.rows[0])
    prevent_row_split(t.rows[0])
    for i, text in enumerate(headers):
        cell = t.rows[0].cells[i]
        shade(cell, LIGHT_GRAY)
        cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run(p, text, bold=True, size=font_size, color=GREEN)
        cell.width = Inches(widths[i])
    for values in rows:
        cells = t.add_row().cells
        prevent_row_split(t.rows[-1])
        for i, value in enumerate(values):
            cell = cells[i]
            cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if i == 0:
                shade(cell, LIGHT_GREEN)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            run(p, str(value), bold=(i == 0), size=font_size)
            cell.width = Inches(widths[i])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    return t


def hyperlink(paragraph, text, url):
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    link_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), GREEN)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    link_run.append(r_pr)
    node = OxmlElement("w:t")
    node.text = text
    link_run.append(node)
    link.append(link_run)
    paragraph._p.append(link)


doc = Document(TEMPLATE)
clear_body(doc)
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

for footer in (section.footer, section.first_page_footer):
    for node in footer._element.iter(qn("w:t")):
        if node.text and "Report Name" in node.text:
            node.text = node.text.replace("Report Name", "Week 3 Homework")

for _ in range(4):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(p, "BIOINFORMATICS & SPATIAL OMICS", bold=True, size=14, color=GREEN)
p.paragraph_format.space_after = Pt(34)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(p, "Week 3 Homework", bold=True, size=34, color=GREEN)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(p, "GEO dataset reconnaissance and paired bulk RNA-seq design", size=16, color="444444")
p.paragraph_format.space_after = Pt(120)
for text, bold, size in (
    ("Ye Gulin", True, 13),
    ("Student ID: SUAT24000191", False, 11),
    ("September 13, 2026", False, 11),
):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, text, bold=bold, size=size)
doc.add_page_break()

heading(doc, "Scope and evidence status")
body(doc, "GSE146853 is used for Part A because it is the IBS host-transcriptome GEO series underlying the IBS cohort discussed in the Week 2 Nature Microbiology paper. The GEO-linked primary study is Mars et al. (Cell, 2020), which is distinct from the Week 2 cross-disease integration paper by Priya et al. (Nature Microbiology, 2022).")
body(doc, "Verification performed: both downloaded gzip files passed integrity checks; matrix dimensions and integer status were checked in R; pairing and IRI status were reconstructed from GEO family metadata; the example GSM-to-SRR chain was checked against the NCBI SRA record.")

heading(doc, "Part A — GSE146853 reconnaissance")
rows_a = [
    ("1. Data-driven question", "Among adults aged 18–65 with IBS-C or IBS-D, compared with matched healthy controls, which colonic mucosal host genes show differential expression in bulk RNA-seq data after accounting for repeated biopsies from the same participant?"),
    ("2. Study identity", "GSE146853 — “Longitudinal multi-omics reveals subset-specific mechanisms underlying irritable bowel syndrome”; Homo sapiens; bulk RNA-seq; GPL16791, Illumina HiSeq 2500. GEO-linked publication: Mars et al., Cell 2020, PMID 32916129. Week 2 paper using this IBS cohort: Priya et al., Nature Microbiology 2022."),
    ("3. Accession map", "GSE146853 → BioProject PRJNA612180 → SRA study SRP252532. Example: GSM4407870 (10007542-1) → SRA experiment SRX7899619 → run SRR11294060; BioSample SAMN14364294."),
    ("4. Available data layers", "GEO provides a Series Matrix, supplementary raw gene-count matrix GSE146853_GeneCount_raw.txt.gz, and paired-end raw reads in SRA. No single-cell files are listed. The wider publication contains other omics, but this GEO series itself is host bulk RNA-seq rather than the entire multi-omics study."),
    ("5. Reconstructed design", "68 biopsy RNA-seq samples from 42 subject IDs: 26 subjects have two biopsies and 16 have one. Raw metadata summarize as IBS-C/Constipation 23 samples from 15 subjects, IBS-D/Diarrhea 23 samples from 13 subjects (including two records with a duplicated-label typo), Healthy 21 samples from 13 subjects, and one sample with cohort recorded as NA. Recorded covariates are age, sex and BMI; controls were matched on these variables. Repeated samples must be modeled by subject. All samples use GPL16791; an explicit sequencing-batch variable is not provided, so batch is unknown rather than assumed absent."),
]
table(doc, ["Step", "Evidence recorded"], rows_a, [1.50, 5.50], 9.0)

heading(doc, "Local matrix verification", 2)
table(
    doc,
    ["Check", "Observed result", "Interpretation"],
    [
        ("File integrity", "gzip test passed", "The manually downloaded file is complete."),
        ("Dimensions", "64,253 gene rows × 68 sample columns", "All 68 GEO sample columns are present."),
        ("Value type", "Finite, non-negative integers", "The file is a raw count matrix suitable for count-based methods after design and QC decisions."),
    ],
    [1.35, 2.20, 3.45],
    8.6,
)

heading(doc, "Part B — GSE87487 analysis decisions")
rows_b = [
    ("Are the values counts?", "Yes. The matrix contains 60,498 gene rows and 20 sample columns. Every sample value is finite, non-negative and integer-valued. GEO metadata state that reads were aligned to hg38 and summarized with featureCounts."),
    ("Are samples independent?", "No. The 20 biopsies are 10 matched subject pairs, each with one pre-reperfusion (Bx1) and one post-reperfusion (Bx2) sample. Subject ID must remain in the metadata and model; treating all columns as independent would be pseudoreplication."),
    ("Can all 20 be compared directly?", "Not as an unpaired 10-versus-10 test and not while silently mixing IRI status. The prespecified contrast here is post-reperfusion versus pre-reperfusion within the four IRI-positive subjects (8 samples), using design ~ subject + stage. The six IRI-negative pairs are outside this contrast."),
    ("Is full raw-read reprocessing possible?", "Yes. Raw reads are available under SRA study SRP090633 (BioProject PRJNA344898). A full workflow could download FASTQ, run QC, align to the stated genome build and recount genes, but this is beyond the one-hour homework and unnecessary for checking the supplied count matrix."),
]
table(doc, ["Question", "Student decision and evidence"], rows_b, [1.80, 5.20], 9.1)

heading(doc, "Design reconstruction for GSE87487", 2)
table(
    doc,
    ["Item", "Reconstructed value"],
    [
        ("Subjects and pairing", "10 subjects; exactly one pre and one post biopsy per subject."),
        ("IRI-positive", "4 pairs: RJ, Pt9, Pt10 and Pt12."),
        ("IRI-negative", "6 pairs: HB, MJ, OC, Pt2, Pt6 and Pt8."),
        ("Selected contrast", "IRI-positive post versus pre; DESeq2 design = ~ subject + stage; contrast = stage post versus pre."),
        ("Boundary", "The task requires a justified design decision, not a claim of clinical validation. Small subgroup size and unreported batch covariates limit inference."),
    ],
    [1.80, 5.20],
    9.1,
)

heading(doc, "Evidence links", 2)
for label, url in (
    ("GSE146853 GEO Series", "https://www.ncbi.nlm.nih.gov/geo/query/acc.cgi?acc=GSE146853"),
    ("GSM4407870 GEO Sample", "https://www.ncbi.nlm.nih.gov/geo/query/acc.cgi?acc=GSM4407870"),
    ("Mars et al. 2020 PubMed record", "https://pubmed.ncbi.nlm.nih.gov/32916129/"),
    ("GSE87487 GEO Series", "https://www.ncbi.nlm.nih.gov/geo/query/acc.cgi?acc=GSE87487"),
):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    run(p, "• ", size=10, color=GREEN)
    hyperlink(p, label, url)

heading(doc, "Appendix — R script used for verification and DESeq2")
body(doc, "This is the same script supplied with the homework. It reads both matrices, verifies counts and pairing, and runs the prespecified paired contrast.", italic=True)
for line in R_SCRIPT.read_text(encoding="utf-8").splitlines():
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 0.88
    run(p, line if line else " ", size=7.2, font="Menlo")

doc.core_properties.title = "Week 3 Homework — GEO Dataset Reconnaissance"
doc.core_properties.author = "Ye Gulin"
doc.core_properties.subject = "GSE146853 and GSE87487"
doc.core_properties.keywords = "GEO, RNA-seq, DESeq2, paired design"
doc.save(OUTPUT)
print(OUTPUT)
